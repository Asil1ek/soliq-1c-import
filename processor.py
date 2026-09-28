import os
import re
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def clean_str(val):
    if pd.isna(val):
        return ""
    val_str = str(val).strip()
    if val_str.endswith(".0"):
        val_str = val_str[:-2]
    return val_str

def format_ikpu(val):
    if pd.isna(val):
        return ""
    s = clean_str(val)
    digits = re.sub(r'\D', '', s)
    if not digits:
        return s
    if len(digits) < 17:
        digits = digits.zfill(17)
    return digits

def format_barcode(val):
    if pd.isna(val):
        return ""
    return clean_str(val)

def load_nomenklatura_mapping(file_paths):
    """
    Loads nomenclature mapping from one or more files.
    Accepts a single path string or a list of paths.
    Merges all files into combined lookup dictionaries.
    Returns dicts: by_barcode, by_ikpu, by_name
    """
    if not file_paths:
        return {}, {}, {}

    if isinstance(file_paths, str):
        file_paths = [file_paths]

    by_barcode = {}
    by_ikpu = {}
    by_name = {}

    for file_path in file_paths:
        if not file_path or not os.path.exists(file_path):
            continue
        try:
            df = pd.read_excel(file_path)
            col_code = None
            col_ikpu = None
            col_barcode = None
            col_name = None

            for c in df.columns:
                cs = str(c).lower()
                if 'код' in cs and not col_code:
                    col_code = c
                elif ('икпу' in cs or 'мхик' in cs) and not col_ikpu:
                    col_ikpu = c
                elif ('штрих' in cs or 'артикул' in cs or 'barcode' in cs) and not col_barcode:
                    col_barcode = c
                elif ('наименование' in cs or 'ном' in cs or 'tovar' in cs) and not col_name:
                    col_name = c

            if col_code:
                for _, row in df.iterrows():
                    code = clean_str(row[col_code])
                    if not code:
                        continue
                    if col_barcode and pd.notna(row.get(col_barcode)):
                        bc = format_barcode(row[col_barcode])
                        if bc and bc not in by_barcode:
                            by_barcode[bc] = code
                    if col_ikpu and pd.notna(row.get(col_ikpu)):
                        ik = format_ikpu(row[col_ikpu])
                        if ik and ik not in by_ikpu:
                            by_ikpu[ik] = code
                    if col_name and pd.notna(row.get(col_name)):
                        nm = clean_str(row[col_name]).lower()
                        if nm and nm not in by_name:
                            by_name[nm] = code

        except Exception as e:
            print(f"Error reading nomenclature file {file_path}: {e}")

    return by_barcode, by_ikpu, by_name


def process_soliq_file(file_path, nomenklatura_mapping=None, default_code=""):
    """
    Processes a Soliq checks file (either Sotuv or Qaytarish).
    Returns:
    - summary: dict of summary numbers
    - df_grouped: grouped by product (Пономенклатурно)
    - df_raw: item lines preserved
    """
    by_barcode, by_ikpu, by_name = nomenklatura_mapping or ({}, {}, {})

    df = pd.read_excel(file_path, header=1)

    # Find relevant columns
    col_tovar = None
    col_miqdor = None
    col_narx = None
    col_qqs = None
    col_naqd = None
    col_karta = None
    col_ikpu = None
    col_barcode = None
    col_unit = None
    col_chek = None
    col_fm = None

    for c in df.columns:
        cs = str(c).strip()
        cl = cs.lower()
        if 'маҳсулот номи' in cl or 'махсулот номи' in cl:
            col_tovar = c
        elif 'миқдори' in cl or 'микдори' in cl:
            col_miqdor = c
        elif cl == 'нархи' or cl.startswith('нархи'):
            col_narx = c
        elif cl == 'ққс суммаси' or cl == 'ккс суммаси' or 'ққс суммаси' in cl:
            col_qqs = c
        elif 'жами нақд' in cl or 'жами накд' in cl:
            col_naqd = c
        elif 'жами банк карта' in cl:
            col_karta = c
        elif 'маҳсулот коди' in cl or 'махсулот коди' in cl:
            col_ikpu = c
        elif 'штрих код' in cl or 'штрихкод' in cl:
            col_barcode = c
        elif 'ўлчов бирлиги' in cl or 'улчов бирлиги' in cl:
            col_unit = c
        elif 'чек рақами' in cl or 'чек раками' in cl:
            col_chek = c
        elif 'фм рақами' in cl or 'фм раками' in cl:
            col_fm = c

    df['TovarNomi'] = df[col_tovar].astype(str).str.strip() if col_tovar else "Товар"
    df['Miqdori'] = pd.to_numeric(df[col_miqdor], errors='coerce').fillna(0.0) if col_miqdor else 1.0
    df['Summa'] = pd.to_numeric(df[col_narx], errors='coerce').fillna(0.0) if col_narx else 0.0
    df['QQS'] = pd.to_numeric(df[col_qqs], errors='coerce').fillna(0.0) if col_qqs else 0.0

    df['IKPU'] = df[col_ikpu].apply(format_ikpu) if col_ikpu else ""
    df['Barcode'] = df[col_barcode].apply(format_barcode) if col_barcode else ""
    df['UnitCode'] = df[col_unit].apply(clean_str) if col_unit else "1352017"

    df['Narxi'] = (df['Summa'] / df['Miqdori']).round(2)
    df['Narxi'] = df['Narxi'].fillna(0.0)
    df['Narxi'] = df['Narxi'].replace([float('inf'), float('-inf')], 0.0)

    # Map nomenclature code — empty string if not found (not default_code)
    def find_code(row):
        bc = row['Barcode']
        ik = row['IKPU']
        nm = row['TovarNomi'].lower()
        if bc and bc in by_barcode:
            return by_barcode[bc]
        if ik and ik in by_ikpu:
            return by_ikpu[ik]
        if nm and nm in by_name:
            return by_name[nm]
        # Kod topilmadi — bo'sh qoldirish
        return ""

    df['Kod'] = df.apply(find_code, axis=1)
    df['StavkaNDS'] = "12%"

    # Receipt-level payment calculations
    receipt_subset = []
    if col_fm: receipt_subset.append(col_fm)
    if col_chek: receipt_subset.append(col_chek)

    if receipt_subset and col_naqd and col_karta:
        receipts = df.drop_duplicates(subset=receipt_subset)
        total_cash = float(pd.to_numeric(receipts[col_naqd], errors='coerce').fillna(0).sum())
        total_card = float(pd.to_numeric(receipts[col_karta], errors='coerce').fillna(0).sum())
    else:
        receipts = df
        total_cash = 0.0
        total_card = 0.0

    total_sum = float(df['Summa'].sum())
    total_vat = float(df['QQS'].sum())

    summary = {
        'total_sum': total_sum,
        'total_cash': total_cash,
        'total_card': total_card,
        'total_vat': total_vat,
        'total_rows': len(df),
        'unique_receipts': len(receipts) if receipt_subset else len(df),
        'unique_products': int(df['TovarNomi'].nunique()),
        'unique_barcodes': int(df['Barcode'].replace('', pd.NA).dropna().nunique()),
        'unique_ikpu': int(df['IKPU'].replace('', pd.NA).dropna().nunique()),
        'matched_codes': int((df['Kod'] != '').sum()),
        'unmatched_codes': int((df['Kod'] == '').sum()),
    }

    # Grouped dataset (Пономенклатурно — grouped by product)
    grouped = df.groupby(['Kod', 'TovarNomi', 'Barcode', 'IKPU', 'StavkaNDS'], dropna=False, as_index=False).agg({
        'Miqdori': 'sum',
        'Summa': 'sum',
        'QQS': 'sum'
    })
    grouped['Narxi'] = (grouped['Summa'] / grouped['Miqdori']).round(2).fillna(0.0)
    grouped['Narxi'] = grouped['Narxi'].replace([float('inf'), float('-inf')], 0.0)

    cols_order = ['Kod', 'TovarNomi', 'Barcode', 'IKPU', 'Miqdori', 'Narxi', 'Summa', 'StavkaNDS', 'QQS']
    df_grouped = grouped[cols_order].sort_values(by='Summa', ascending=False).reset_index(drop=True)
    df_raw = df[cols_order].reset_index(drop=True)

    return summary, df_grouped, df_raw


def subtract_returns_from_sales(sotuv_grp, qayt_grp):
    """
    Sotuvdan qaytarishlarni ayirish.
    Bir xil tovarni (TovarNomi + Barcode + IKPU) bo'yicha topib,
    miqdor va summani kamaytiradi.
    Agar tovar butunlay qaytarilgan bo'lsa — ro'yxatdan o'chiriladi.
    """
    if qayt_grp is None or qayt_grp.empty:
        return sotuv_grp.copy()

    # Sotuv nusxasi
    net = sotuv_grp.copy()

    # Qaytarish bo'yicha iteratsiya
    for _, qrow in qayt_grp.iterrows():
        # Bir xil tovarni topish
        mask = (
            (net['TovarNomi'] == qrow['TovarNomi']) &
            (net['Barcode'] == qrow['Barcode']) &
            (net['IKPU'] == qrow['IKPU'])
        )
        idx = net.index[mask]
        if len(idx) > 0:
            i = idx[0]
            net.at[i, 'Miqdori'] = net.at[i, 'Miqdori'] - qrow['Miqdori']
            net.at[i, 'Summa'] = net.at[i, 'Summa'] - qrow['Summa']
            net.at[i, 'QQS'] = net.at[i, 'QQS'] - qrow['QQS']

    # Miqdori <= 0 bo'lganlarni o'chirish
    net = net[net['Miqdori'] > 0].reset_index(drop=True)

    # Narxni qayta hisoblash
    net['Narxi'] = (net['Summa'] / net['Miqdori']).round(2).fillna(0.0)
    net['Narxi'] = net['Narxi'].replace([float('inf'), float('-inf')], 0.0)

    return net


def export_to_1c_excel(df_data, output_path, title="1C Загрузка"):
    """
    Exports DataFrame to an Excel file formatted for 1C import.
    Columns: Код номенклатуры, Номенклатура, Артикул, ИКПУ,
             Количество, Цена, Сумма, Ставка НДС, Сумма НДС
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Данные для 1С"

    headers = [
        "Код номенклатуры",
        "Номенклатура",
        "Артикул",
        "ИКПУ",
        "Количество",
        "Цена",
        "Сумма",
        "Ставка НДС",
        "Сумма НДС"
    ]

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    ws.row_dimensions[1].height = 26
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border

    row_num = 2
    for _, row in df_data.iterrows():
        ws.row_dimensions[row_num].height = 20

        c1 = ws.cell(row=row_num, column=1, value=str(row['Kod']) if row['Kod'] else "")
        c1.alignment = align_center
        c1.border = thin_border
        c1.number_format = '@'

        c2 = ws.cell(row=row_num, column=2, value=str(row['TovarNomi']))
        c2.alignment = align_left
        c2.border = thin_border

        c3 = ws.cell(row=row_num, column=3, value=str(row['Barcode']) if row['Barcode'] else "")
        c3.alignment = align_center
        c3.border = thin_border
        c3.number_format = '@'

        c4 = ws.cell(row=row_num, column=4, value=str(row['IKPU']) if row['IKPU'] else "")
        c4.alignment = align_center
        c4.border = thin_border
        c4.number_format = '@'

        c5 = ws.cell(row=row_num, column=5, value=float(row['Miqdori']))
        c5.alignment = align_right
        c5.border = thin_border
        c5.number_format = '#,##0.000'

        c6 = ws.cell(row=row_num, column=6, value=float(row['Narxi']))
        c6.alignment = align_right
        c6.border = thin_border
        c6.number_format = '#,##0.00'

        c7 = ws.cell(row=row_num, column=7, value=float(row['Summa']))
        c7.alignment = align_right
        c7.border = thin_border
        c7.number_format = '#,##0.00'

        c8 = ws.cell(row=row_num, column=8, value=str(row['StavkaNDS']))
        c8.alignment = align_center
        c8.border = thin_border

        c9 = ws.cell(row=row_num, column=9, value=float(row['QQS']))
        c9.alignment = align_right
        c9.border = thin_border
        c9.number_format = '#,##0.00'

        row_num += 1

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 42
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 22

    wb.save(output_path)
    return output_path
